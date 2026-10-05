import logging
import os
import time
from fastapi import APIRouter, File, UploadFile, HTTPException, Depends, BackgroundTasks
from fastapi.concurrency import run_in_threadpool
from typing import List
from .schemas import QueryRequest, QueryResponse, StatusResponse
from .auth import get_current_user
from core.text_utils import clean_spaced_text
from storage.ingestion import ingestion_pipeline
from storage.vector_db import vector_db
from storage.memory_cache import memory_cache
from storage.graph_db import graph_db
from agents.orchestrator import process_query_workflow
from core.embeddings import LocalEmbedder
from immune.macrophage import immune_system
from core.confluence import confluence_sync

logger = logging.getLogger(__name__)

api_router = APIRouter()

# Documents under this many extracted characters bypass embedding entirely.
DIRECT_INJECTION_THRESHOLD = 15000

# Uploads are read fully into memory before extraction, so cap them.
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_MB", "25")) * 1024 * 1024


def process_upload_in_background(filename: str, text_content: str, user_id: str):
    try:
        # Re-uploading a filename replaces it: without this, the old chunks stay in
        # Qdrant forever and the document is silently double-counted at retrieval.
        removed = vector_db.delete_by_filename(filename, owner_id=user_id)
        if removed:
            logger.info(f"Replacing '{filename}': removed {removed} stale vectors before re-indexing.")
        graph_db.delete_by_filename(filename, owner_id=user_id)

        nodes = ingestion_pipeline.process_document(filename, text_content)
        embedder = LocalEmbedder.get_embedder()
        texts = [n["text"] for n in nodes]
        real_vectors = embedder.embed_documents(texts)
        payloads = [
            {"text": n["text"], "file_name": filename, "user_id": str(user_id)}
            for n in nodes
        ]
        vector_db.upsert_vectors(real_vectors, payloads)
        # Populate graph DB with keyword relationships from this document
        ingestion_pipeline.populate_graph(filename, nodes, owner_id=user_id)
        logger.info(f"Background processing complete for {filename}: {len(nodes)} chunks indexed.")

        # Large documents are exactly the ones most likely to contain the
        # contradictions the auditor exists to catch, so scan them too.
        immune_system.scan_for_conflicts(text_content, filename, owner_id=user_id)
    except Exception:
        logger.exception(f"Background upload processing failed for {filename}")


@api_router.post("/upload", response_model=StatusResponse)
async def upload_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user)
):
    """
    Ingests PDF, TXT, DOCX, PPTX with proper text extraction per file type.
    """
    owner_id = current_user["id"]
    try:
        content = await file.read()
        if len(content) > MAX_UPLOAD_BYTES:
            raise HTTPException(
                status_code=413,
                detail=f"File exceeds the {MAX_UPLOAD_BYTES // (1024 * 1024)}MB upload limit.",
            )
        filename = file.filename
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "txt"

        # --- Proper text extraction per file format ---
        text_content = ""
        if ext == "csv":
            try:
                import io, csv
                decoded = content.decode("utf-8-sig", errors="ignore")
                reader = csv.reader(io.StringIO(decoded))
                headers = next(reader, None)
                rows_text = []
                for row_idx, row in enumerate(reader):
                    row_details = []
                    for i, val in enumerate(row):
                        if val.strip():
                            header = headers[i] if headers and i < len(headers) else f"Column {i+1}"
                            row_details.append(f"{header}: {val.strip()}")
                    if row_details:
                        rows_text.append(f"[Record {row_idx+1}] " + " | ".join(row_details))
                text_content = "\n\n".join(rows_text)
                logger.info(f"Extracted {len(rows_text)} records from CSV '{filename}'")
            except Exception:
                logger.exception("CSV extraction failed")
                raise HTTPException(status_code=422, detail="Could not read this CSV file.")

        elif ext == "pptx":
            try:
                import io
                from pptx import Presentation
                prs = Presentation(io.BytesIO(content))
                slide_texts = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_texts.append(shape.text.strip())
                text_content = "\n\n".join(slide_texts)
                logger.info(f"Extracted {len(slide_texts)} text blocks from PPTX '{filename}'")
            except Exception:
                logger.exception("PPTX extraction failed")
                raise HTTPException(status_code=422, detail="Could not read this PPTX file.")

        elif ext == "pdf":
            try:
                import io
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(content))
                pages = [page.extract_text() or "" for page in reader.pages]
                raw_text = "\n\n".join(p for p in pages if p.strip())
                text_content = clean_spaced_text(raw_text)
                logger.info(f"Extracted and Sanitized {len(reader.pages)} pages from PDF '{filename}'")
            except Exception:
                logger.exception("PDF extraction failed")
                raise HTTPException(status_code=422, detail="Could not read this PDF file.")

        elif ext == "docx":
            try:
                import io, docx
                doc = docx.Document(io.BytesIO(content))
                text_content = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
                logger.info(f"Extracted {len(doc.paragraphs)} paragraphs from DOCX '{filename}'")
            except Exception as e:
                logger.warning(f"DOCX extraction failed, falling back to UTF-8: {e}")
                text_content = content.decode("utf-8", errors="ignore")

        elif ext in ["png", "jpg", "jpeg"]:
            try:
                import base64
                from langchain_google_genai import ChatGoogleGenerativeAI
                from langchain_core.messages import HumanMessage

                api_key = os.getenv("GOOGLE_API_KEY") or os.getenv("GEMINI_API_KEY")
                if not api_key or api_key.startswith("paste-") or api_key.startswith("your-"):
                    raise HTTPException(
                        status_code=503,
                        detail="Vision OCR is unavailable: Valid GOOGLE_API_KEY is not configured.",
                    )

                base64_image = base64.b64encode(content).decode('utf-8')
                mime_type = f"image/{'jpeg' if ext == 'jpg' else ext}"

                llm = ChatGoogleGenerativeAI(model="gemini-2.5-pro", temperature=0.0, google_api_key=api_key)

                prompt = (
                    "You are an elite Industrial OCR System. "
                    "Extract all text, labels, measurements, component names, and structural data from this engineering diagram or scanned form. "
                    "Format the output as a clear, highly structured markdown document. Do not miss any numbers or technical specifications."
                )

                # Vision calls are synchronous and slow; keep them off the event loop.
                msg = await run_in_threadpool(
                    llm.invoke,
                    [HumanMessage(content=[
                        {"type": "text", "text": prompt},
                        {"type": "image_url", "image_url": f"data:{mime_type};base64,{base64_image}"}
                    ])],
                )

                text_content = f"--- Vision OCR Extraction for {filename} ---\n\n" + msg.content
                logger.info(f"Vision OCR completed successfully for '{filename}'")
            except HTTPException:
                raise
            except Exception:
                logger.exception("Image Vision extraction failed")
                raise HTTPException(status_code=422, detail="Could not analyze this image.")

        else:  # txt, md, etc.
            text_content = content.decode("utf-8", errors="ignore")

        if not text_content.strip():
            raise HTTPException(status_code=422, detail="No readable text could be extracted from this file.")

        file_size = len(text_content)
        logger.info(f"Received file: {filename} (Extracted {file_size} chars) from {current_user['username']}")

        # Small file: store in memory cache directly (bypass RAG)
        if file_size < DIRECT_INJECTION_THRESHOLD:
            graph_db.delete_by_filename(filename, owner_id=owner_id)  # replace on re-upload
            memory_cache.store_file(filename, text_content, owner_id=owner_id)
            # Create nodes for graph indexing even if bypassing vector search
            temp_nodes = [{"text": text_content, "metadata": {"file_name": filename}}]
            ingestion_pipeline.populate_graph(filename, temp_nodes, owner_id=owner_id)

            background_tasks.add_task(immune_system.scan_for_conflicts, text_content, filename, owner_id)
            return StatusResponse(
                status="success",
                message=f"Document '{filename}' saved to Knowledge Graph. Macrophage scan initiated.",
                data={"chunks_created": 1, "status": "direct_injection"}
            )

        # Large file: embed in background
        background_tasks.add_task(process_upload_in_background, filename, text_content, owner_id)
        return StatusResponse(
            status="success",
            message=f"Document '{filename}' successfully ingested. Embedding processing running in background.",
            data={"status": "processing"}
        )
    except HTTPException:
        raise
    except Exception:
        logger.exception("Error uploading file")
        raise HTTPException(status_code=500, detail="Upload failed. See server logs for details.")


@api_router.post("/query", response_model=QueryResponse)
async def query_knowledge(
    request: QueryRequest,
    current_user: dict = Depends(get_current_user)
):
    """
    Handles natural language queries and triggers the Multi-Agent LangGraph workflow.
    """
    try:
        logger.info(f"Received query: {request.query} in mode: {request.mode} for user: {current_user['username']}")

        # The LangGraph run is fully synchronous and takes seconds — running it
        # inline would block the event loop for every other request.
        start_time = time.time()
        agent_raw_result = await run_in_threadpool(
            process_query_workflow,
            request.query,
            mode=request.mode,
            files=request.files,
            owner_id=current_user["id"],
        )
        end_time = time.time()

        logger.info(f"LangGraph execution completed in {round(end_time-start_time, 2)}s.")

        # Extract metadata from state
        citations = agent_raw_result.get("citations", [])
        sources = list(set([c.get("source", "Unknown") for c in citations if c.get("source")]))

        # Mode-aware source labeling
        if not sources:
            if request.mode == "Online":
                sources = ["No Web Results Found"]
            elif request.mode == "Hybrid":
                sources = ["No matches in Files or Web"]
            else:
                sources = ["No matches in Private Files"]
        elif request.mode == "Online":
            # Filter out any internal labels that might have leaked
            sources = [s for s in sources if "Web:" in s]
            if not sources:
                sources = ["Web (Direct Fetch)"]

        safe_citations = [
            {
                "source": c.get("source", "Unknown"),
                "snippet": (c.get("content") or "")[:400],
                "url": c.get("url")
            }
            for c in citations
            if isinstance(c, dict)
        ]

        return QueryResponse(
            answer=agent_raw_result.get("final_answer", "No answer could be generated."),
            confidence_score=int(agent_raw_result.get("confidence", 0.0)),
            confidence_level="High" if agent_raw_result.get("confidence", 0) > 85 else "Medium",
            strategy=agent_raw_result.get("strategy", "Strict Hybrid Validation"),
            sources=sources,
            citations=safe_citations,
            steps=agent_raw_result.get("steps_taken", [])
        )
    except Exception:
        logger.exception("Error executing query")
        raise HTTPException(status_code=500, detail="Query failed. See server logs for details.")


@api_router.delete("/documents/{filename}", response_model=StatusResponse)
async def delete_document(
    filename: str,
    current_user: dict = Depends(get_current_user)
):
    """
    Deletes a document from the vector DB, memory cache and graph — this user's copy only.
    """
    owner_id = current_user["id"]
    try:
        logger.info(f"User {current_user['username']} requested deletion of {filename}")

        vectors_removed = vector_db.delete_by_filename(filename, owner_id=owner_id)
        cache_removed = memory_cache.delete_file(filename, owner_id=owner_id)
        edges_removed = graph_db.delete_by_filename(filename, owner_id=owner_id)

        if not (vectors_removed or cache_removed or edges_removed):
            # Either it never existed or it belongs to someone else — same answer either way.
            raise HTTPException(status_code=404, detail=f"No document named '{filename}' found.")

        return StatusResponse(
            status="success",
            message=f"Document '{filename}' successfully deleted from databases.",
            data={"deleted_file": filename}
        )
    except HTTPException:
        raise
    except Exception:
        logger.exception(f"Error deleting file {filename}")
        raise HTTPException(status_code=500, detail="Delete failed. See server logs for details.")


@api_router.get("/alerts")
async def get_alerts(current_user: dict = Depends(get_current_user)):
    """Exposes current proactive alerts and patterns to the frontend"""
    alerts = immune_system.list_alerts(owner_id=current_user["id"])
    return {
        "status": "alert" if len(alerts) > 0 else "secure",
        "active_threats": len(alerts),
        "reports": alerts
    }


@api_router.delete("/alerts/{alert_id}")
async def dismiss_alert(alert_id: str, current_user: dict = Depends(get_current_user)):
    """Dismiss a specific alert"""
    if not immune_system.dismiss(alert_id, owner_id=current_user["id"]):
        raise HTTPException(status_code=404, detail="Alert not found.")
    return {"status": "success"}


@api_router.get("/documents")
async def list_documents(current_user: dict = Depends(get_current_user)):
    """
    Returns a combined list of this user's indexed filenames from MemoryCache + VectorDB.
    Used by the frontend to restore the file list after login.
    """
    owner_id = current_user["id"]
    try:
        memory_files = memory_cache.list_files(owner_id=owner_id)
        vector_files = await run_in_threadpool(vector_db.list_filenames, owner_id)

        # Combine and deduplicate, preserving order
        all_files = list(dict.fromkeys(memory_files + vector_files))
        return {"files": all_files}
    except Exception:
        logger.exception("Error listing documents")
        raise HTTPException(status_code=500, detail="Could not list documents.")


@api_router.get("/graph/data")
async def get_graph_data(current_user: dict = Depends(get_current_user)):
    """Exports graph data for the 2D visualizer, scoped to this user."""
    return graph_db.get_graph_data(owner_id=current_user["id"])


@api_router.get("/suggestions")
async def get_dynamic_suggestions(current_user: dict = Depends(get_current_user)):
    """Fetches random graph entities and dynamically templates them into suggested questions"""
    entities = graph_db.get_random_entities(5, owner_id=current_user["id"])

    if not entities:
        return {"suggestions": [
            "What documents are currently uploaded?",
            "Can you summarize the main concepts?",
            "What security policies are available?"
        ]}

    import random
    templates = [
        "Explain the significance of '{0}' in the documents.",
        "What are the main findings regarding '{0}'?",
        "Can you summarize the context around '{0}'?",
        "How is '{0}' related to other concepts?"
    ]

    suggestions = [tmpl.format(entity) for entity, tmpl in zip(entities, [random.choice(templates) for _ in entities])]
    return {"suggestions": suggestions}


@api_router.post("/confluence/sync", response_model=StatusResponse)
async def sync_confluence(
    request: dict,
    current_user: dict = Depends(get_current_user)
):
    domain = request.get("domain")
    space_key = request.get("spaceKey")
    email = request.get("email")
    api_token = request.get("apiToken")

    if not all([domain, space_key, email, api_token]):
        raise HTTPException(status_code=400, detail="Missing Confluence configuration fields")

    # Run sync (this might take a minute depending on space size, client handles timeout or we run it async)
    result = await confluence_sync.sync_space(domain, space_key, email, api_token, current_user['id'])

    if result["status"] == "error":
        logger.error(f"Confluence sync failed: {result.get('message')}")
        raise HTTPException(status_code=502, detail="Confluence sync failed. See server logs for details.")

    return StatusResponse(status="success", message=f"Ingested {result['title']} ({result.get('pages', 0)} pages synced)")


@api_router.get("/graph/stats")
async def get_graph_stats(current_user: dict = Depends(get_current_user)):
    """Debug endpoint to inspect this user's slice of the knowledge graph."""
    owner_id = current_user["id"]
    visible_nodes = graph_db._visible_node_names(owner_id)
    edges = [
        (u, v, d) for u, v, d in graph_db.graph.edges(data=True)
        if graph_db._visible(d, owner_id)
    ]
    my_files = memory_cache.list_files(owner_id=owner_id)
    return {
        "total_nodes": len(visible_nodes),
        "total_edges": len(edges),
        "sample_nodes": list(visible_nodes)[:20],
        "sample_edges": [{"from": u, "to": v, "rel": d.get("relation", "?")} for u, v, d in edges[:20]],
        "memory_cache_files": my_files,
        "memory_cache_total_chars": sum(
            len(memory_cache.cache[f]["content"]) for f in my_files
        ),
    }
